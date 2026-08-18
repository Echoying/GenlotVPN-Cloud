#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把分享文案填进《追逐梦想-科技智享》原片。

按版式选页，不按页序硬塞：箭头页放流程，循环页放共创，
递减圆放评分，拼图放四块架构，问号页放边界。

运行：python pipeline/work/_shared/build_share_pptx.py
"""
from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

HERE = Path(__file__).resolve().parent
OUT = HERE / "分享-超级AI个体-40分钟.pptx"
TPL = HERE / "share-assets" / "tpl" / "追逐梦想-科技智享.pptx"


def set_text(shape, text):
    if shape is None or not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.word_wrap = True
    lines = text.split("\n") if isinstance(text, str) else list(text)
    paras = list(tf.paragraphs)
    for i, para in enumerate(paras):
        line = lines[i] if i < len(lines) else ""
        if para.runs:
            para.runs[0].text = line
            for run in para.runs[1:]:
                run.text = ""
        elif line:
            para.add_run().text = line
    if len(lines) > len(paras) and paras:
        paras[0].runs[0].text = "\n".join(lines) if paras[0].runs else text


def sh(slide, *idx):
    cur = slide.shapes
    obj = None
    for i in idx:
        obj = cur[i]
        cur = obj.shapes if obj.shape_type == MSO_SHAPE_TYPE.GROUP else None
    return obj


def fill_slide(slide, mapping):
    for key, text in mapping.items():
        if isinstance(key, int):
            set_text(slide.shapes[key], text)
        else:
            set_text(sh(slide, *key), text)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def build():
    if not TPL.is_file():
        raise SystemExit("缺少模板：{}".format(TPL))
    tmp = OUT.with_name("_tpl_fill.pptx")
    shutil.copyfile(TPL, tmp)
    prs = Presentation(str(tmp))
    s = prs.slides

    # 1 封面：只换字，半闭框和大字构图不动
    fill_slide(s[0], {
        2: "内部汇报",
        6: "一人做统一 VPN\n六阶段管住交付",
        12: "GENLOT VPN",
        13: "超",
        14: "级",
        15: "个",
        16: "体",
        17: "一人交付",
    })
    notes(s[0], "开场：不是 AI 写了个 VPN，也不是人选完栈 AI 只写码。两条主线：共创选型 + 六阶段轨道。")

    # 2 目录（模板 01 在最下、形状顺序是 02/03/04/01）
    fill_slide(s[1], {
        21: "需求与共创",
        22: "五条目标，人给约束，AI 列案",
        9: "架构与轨道",
        10: "若依 · Qt · TLS · 六阶段",
        13: "两天证据",
        14: "08-15 线路权限 · 08-16 版本拦截",
        17: "评分与边界",
        18: "价值 35 · 效率 25 · 深度 15 · 沉淀合规 25",
    })
    notes(s[1], "入选 60，推荐 70。深度靠六阶段正文。效率不写倍数。")

    # 3 节首页：四个短标题，放「当时要交的事」
    fill_slide(s[2], {
        4: "需求与共创",
        5: "WHAT THE BUSINESS ASKED",
        6: "统一登录\n三级授权",
        7: "双端可用\n多线路同步",
    })
    notes(s[2], "来自 SOP §5.3。审计看板口头补一句。难处在两套 ID，不在页面。")

    # 4 左三右一+灯泡：能力表，S0 单独拎出来
    fill_slide(s[3], {
        14: "当时一个人会什么",
        15: "ONE PERSON, FIVE DOMAINS",
        7: "Java / 若依",
        6: "熟悉。对照已有 Service 写，人抽查接口。",
        9: "管理端 Vue",
        8: "当初陌生。对照用户页来长，人批路径和权限字。",
        11: "Qt 桌面",
        10: "当初陌生。先最小 Demo，传输层不批量生成。",
        13: "S0 传输",
        12: "只改配置和文档。PHASE2 人勾，不自动验收。",
    })
    notes(s[3], "约束来自这张表。易安联同步口头补：一般，人审映射。")

    # 5 四列带箭头：只放顺序过程 → 六阶段收成四步
    fill_slide(s[4], {
        21: "六阶段，同一套",
        22: "ONE TRACK, FOUR MOVES",
        7: "拆解",
        17: "口述收成表。人认边界，标熟悉度。",
        9: "批契约",
        15: "写成 spec。人批准一次。未批不写码。",
        8: "分流",
        16: "陌生先 Demo。熟悉对照已有代码写。",
        18: "门禁",
        20: "gate → deploy → accept。回修最多 2 轮。",
    })
    notes(s[4], "深度分的正文。客户端不进 accept。超 2 轮停下来等人。")

    # 6 齿轮：多方案决策，不是清单
    fill_slide(s[5], {
        9: "列过案，人拍板",
        10: "AI LISTS, HUMAN DECIDES",
        12: "后台",
        11: "列过自研和升 Boot 3。取若依 3.6.8，业务只进 yianlian。",
        14: "否 Web",
        13: "先做了用户端网页。人抓包，看见密码，否了。",
        16: "否 Electron",
        15: "出包快，但按安全和 C++ 约束否了。",
        18: "取 Qt + TLS",
        17: "熟 C++，双端。桌面走 9443，不复用网关 JSON。",
        20: "硬门槛",
        19: "没批准 spec，不写业务码。",
    })
    notes(s[5], "前半场核心。AI 把方案铺开，人按约束否决。若依只是管理壳。")

    # 7 节 02：四块架构预告
    fill_slide(s[6], {
        4: "架构与轨道",
        5: "TWO SURFACES, ONE TRACK",
        6: "管理端 HTTP\n桌面 TLS 9443",
        7: "本地 30303\n两套 ID 映射",
    })
    notes(s[6], "三块说清楚即可。不要现场讲 Feign。")

    # 8 拼图：四块拼成一套系统
    fill_slide(s[7], {
        12: "现在的四块",
        13: "FOUR PIECES, ONE SYSTEM",
        8: "管理端",
        9: "桌面",
        10: "隧道",
        11: "映射",
        15: "管理端",
        14: "浏览器到网关到 yianlian，走 HTTP。权限菜单不另造。",
        17: "桌面",
        16: "直连 vpn-auth。TLS 1.3 + Protobuf，端口 9443。",
        19: "本地 30303",
        18: "仍是 HTTP。厂商控制器，改不了。",
        21: "两套 ID",
        20: "本地改一条，要同步到多台易安联。",
    })
    notes(s[7], "真正难的是 *YianlianMapping。两端共用 proto/vpn，契约人确认。")

    # 9 四箭头循环+对话泡：共创闭环
    fill_slide(s[8], {
        4: "定框架时转这一圈",
        5: "CONSTRAINT  ->  OPTIONS  ->  DECISION",
        9: "人给约束",
        8: "一人、Java 8 / C++、必须 Win+Mac、必须过抓包、30303 不能改。",
        13: "AI 列 2～3 案",
        12: "怎么切模块、用哪套栈、风险。不要只问「怎么做」。",
        11: "人拍板落盘",
        10: "选哪个、为何不选其他，写成 ADR / spec。",
        7: "未批不写码",
        6: "没有人批准的 spec，不写业务码。这是硬门槛。",
    })
    notes(s[8], "SOP §6。下一页把传输这一刀单独说。")

    # 10 灯泡+左侧主论+右侧两点：传输归人验
    fill_slide(s[9], {
        27: "传输这一刀，归人验",
        28: "S0 IS HUMAN ONLY",
        26: "否明文 JSON",
        25: "桌面复用网关 HTTP，抓包能看见密码。人否了。取 TLS 1.3 + Protobuf :9443。",
        22: "S0 只人",
        21: "TLS、证书 Pin、密钥不进日志。栈配好了，不等于每条新功能自动过。",
        24: "S1 / S2",
        23: "S1 人审逻辑。S2 脚本可断言。截图对照不进 accept.py。",
    })
    notes(s[9], "没改证书不必重勾整份 S0。不要说客户端已自动验收。")

    # 11 节 03
    fill_slide(s[10], {
        4: "两天证据",
        5: "OPEN THE RUN FOLDER",
        6: "08-15 查询\n回修 2 轮",
        7: "08-16 三层\n对照≠PHASE2",
    })
    notes(s[10], "现场优先打开 2026-08-16 当次目录。不要现场填验证码。")

    # 12 树状四枝：一条功能长出轨道
    fill_slide(s[11], {
        1: "08-15 线路权限",
        2: "ONE FEATURE, THE TRACK GREW",
        8: "只读查询",
        7: "更多 → 查看权限，或操作列查看权限。",
        4: "跑门禁",
        3: "gate → deploy → accept，回修 2 轮过。",
        10: "轨道还软",
        9: "bat 误报、deploy 递归、匿名卷挡住 jar。",
        6: "写进 ADR-003",
        5: "根因进脚本，不靠下次有人记得。",
    })
    notes(s[11], "第一条交功能和轨道。弹窗标题是「线路权限 - {用户}」。")

    # 13 递减圆+百分数位：只放评分，从大到小
    fill_slide(s[12], {
        20: "对照评分表",
        21: "WEIGHTS, NOT MULTIPLIERS",
        9: "35",
        8: "25",
        10: "15",
        11: "10",
        12: "业务价值。统一登录、多线路，是领导要的事。",
        14: "效率提升。一人交三层，周期按日计。",
        13: "应用深度。六阶段写进岗位，不是多开聊天。",
        15: "合规协同。未批不写码，S0 人验。沉淀 15 分下一页讲。",
    })
    notes(s[12], "入选 60，推荐 70。效率不写倍数。沉淀并进留下的那页。")

    # 14 四等圆时间轴：08-16 当天
    fill_slide(s[13], {
        17: "08-16 当天节点",
        18: "TIMESTAMPS, NO MULTIPLIERS",
        5: "11:14",
        6: "12:03",
        7: "12:22",
        8: "13:21",
        9: "gate 全量。服务端构建。",
        10: "deploy vpn。带上 93 vpn-auth。",
        11: "accept 通过。管理端先卡验证码。",
        12: "对照升级窗。客户端只对照，不等于 PHASE2。",
    })
    notes(s[13], "第一次三层都动到。不拦 SDK / 30303。")

    # 15 节 04
    fill_slide(s[14], {
        4: "评分与边界",
        5: "WHAT WE CLAIM, WHAT WE DO NOT",
        6: "对得上的数\n不写倍数",
        7: "未批不写码\nS0 人验",
    })
    notes(s[14], "请领导认就是一个人。")

    # 16 中心装饰 + 四角
    fill_slide(s[15], {
        20: "效率只报对得上的数",
        21: "HEADCOUNT AND TIMESTAMPS",
        19: "编制",
        18: "以前后端 + 前端 + 桌面 + 测试。这次一个人交三层。",
        15: "周期",
        14: "同类小功能按周联调。有 spec 当天管理端 accept。",
        17: "15 日",
        16: "轨道还软，回修 2 轮。查询，并修 bat / 卷 / 断言。",
        13: "16 日",
        12: "同一套轨道。11:14–13:21 三层对照。",
        11: "一人交三层",
    })
    notes(s[15], "不写快了几倍，也不写省了多少万。")

    # 17 中心装饰 + 四条
    fill_slide(s[16], {
        11: "留下的，和不能破的",
        12: "WHAT STAYS",
        10: "SOP / 手册",
        9: "方法可学、可搬。同事两行目标就能开干。",
        6: "ADR + 当次",
        5: "ADR-001 到 004。清单、spec、门禁、截图都在目录里。",
        8: "未批不写码",
        7: "硬门槛。没有人批准的 spec，不写业务码。",
        3: "S0 人验",
        2: "不用页面文案代替抓包。凭证只走本机 envs.yaml。",
        1: "可核验",
    })
    notes(s[16], "沉淀 15、合规 10。真凭证不入库。")

    # 18 四瓣 + 四卡片
    fill_slide(s[17], {
        0: "今天不说的",
        14: "PLEASE CONFIRM THREE THINGS",
        6: "01",
        7: "02",
        8: "03",
        9: "04",
        (12, 1): "不说自动验收",
        (12, 0): "不说无人审合入，也不说客户端已经自动验收。",
        (10, 1): "不说已投产",
        (10, 0): "不说规模化投产，也不说 AI 识验证码回填。",
        (11, 1): "请认任务本身",
        (11, 0): "统一登录和多线路就是任务。少了一套不安全的网页入口。",
        (13, 1): "请认一个人",
        (13, 0): "一个人能交三层，靠的是轨道，不是靠堆人。",
    })
    notes(s[17], "问答：套若依？壳子。密钥？envs。不自动验客户端？S0。")

    # 19 结束
    fill_slide(s[18], {
        2: "内部汇报",
        6: "套若依？壳子。\n密钥走本机 envs。",
        12: "GENLOT VPN",
        17: "可以提问",
    })
    notes(s[18], "留 5 分钟问答。客户端为什么不自动验？S0 不能用文案代替抓包。")

    dest = OUT
    try:
        prs.save(OUT)
    except PermissionError:
        dest = OUT.with_name("分享-超级AI个体-40分钟-新.pptx")
        prs.save(dest)
    tmp.unlink(missing_ok=True)
    print("wrote", dest, "slides", len(prs.slides))


if __name__ == "__main__":
    build()
