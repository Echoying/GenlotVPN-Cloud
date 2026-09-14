#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""「研发+AI」全栈交付工作范式 V2

强调：
1. 从0开始到选择若依底座的决策过程（token成本）
2. Web端到客户端的演进（安全驱动）
3. AI辅助架构设计（跨平台、国际化）
4. 六阶段工作法的沉淀与多项目推广验证

运行：python pipeline/work/_shared/build_share_pptx_v2.py
"""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

HERE = Path(__file__).resolve().parent
OUT = HERE / "分享-研发AI全栈交付-40分钟.pptx"
TPL = HERE / "share-assets" / "tpl" / "追逐梦想-科技智享.pptx"
BAK = HERE / "share-ppt-bak"


def backup_current():
    """覆盖前把现片拷到 share-ppt-bak/，带时间戳。"""
    if not OUT.is_file():
        return None
    BAK.mkdir(parents=True, exist_ok=True)
    dest = BAK / "分享-研发AI全栈交付-40分钟-{}.pptx".format(
        datetime.now().strftime("%Y%m%d-%H%M%S")
    )
    shutil.copyfile(OUT, dest)
    print("backup", dest)
    return dest


def set_text(shape, text):
    if shape is None or not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.word_wrap = True
    lines = text.split("\n") if isinstance(text, str) else list(text)
    paras = list(tf.paragraphs)
    for i, para in enumerate(paras):
        line = lines[i] if i < len(lines) else ""
        if para.runs:
            para.runs[0].text = line
            for run in para.runs[1:]:
                run.text = ""
        elif line:
            para.add_run().text = line
    if len(lines) > len(paras) and paras:
        paras[0].runs[0].text = "\n".join(lines) if paras[0].runs else text


def sh(slide, *idx):
    cur = slide.shapes
    obj = None
    for i in idx:
        obj = cur[i]
        cur = obj.shapes if obj.shape_type == MSO_SHAPE_TYPE.GROUP else None
    return obj


def fill_slide(slide, mapping):
    for key, text in mapping.items():
        if isinstance(key, int):
            set_text(slide.shapes[key], text)
        else:
            set_text(sh(slide, *key), text)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def build():
    if not TPL.is_file():
        raise SystemExit("缺少模板：{}".format(TPL))
    backup_current()
    tmp = OUT.with_name("_tpl_fill_v2.pptx")
    shutil.copyfile(TPL, tmp)
    prs = Presentation(str(tmp))
    s = prs.slides

    # 1 封面
    fill_slide(s[0], {
        2: "技术分享",
        6: "后端+AI完成全栈交付\n架构决策·方法沉淀·多项目验证",
        12: "GENLOT VPN",
        13: "全",
        14: "栈",
        15: "交",
        16: "付",
        17: "AI协作范式",
    })
    notes(s[0], "核心：从0开始的架构决策过程 + 可推广的六阶段工作法 + 多项目验证。")

    # 2 目录
    fill_slide(s[1], {
        21: "架构演进历程",
        22: "从0到若依底座·Web到客户端·token成本驱动",
        9: "人机协作决策",
        10: "AI辅助选型·Qt跨平台·国际化设计",
        13: "六阶段工作法",
        14: "方法沉淀·质量门禁·多项目推广",
        17: "实践与验证",
        18: "4个项目验证·真实数据·可核验证据",
    })
    notes(s[1], "重点：架构决策过程（为什么选若依/Qt）+ 六阶段方法在多项目的验证。")

    # 3 节首页：架构演进
    fill_slide(s[2], {
        4: "架构演进历程",
        5: "FROM ZERO TO PRODUCTION",
        6: "从0开始\ntoken成本高",
        7: "底座选型\nWeb到客户端",
    })
    notes(s[2], "这一段讲清楚「为什么不从0造轮子」「为什么不用Web端」，体现真实决策过程。")

    # 4 第一次尝试：从0开始的问题
    fill_slide(s[3], {
        14: "第一次尝试：从0构建",
        15: "WHY NOT FROM SCRATCH",
        7: "初始方案",
        6: "Spring Boot微服务从0搭建·AI生成基础框架",
        9: "遇到问题",
        8: "Token消耗巨大·基础设施重复造轮子·权限菜单耗时多",
        11: "关键认知",
        10: "基础框架不是核心价值·应聚焦VPN业务逻辑",
        13: "决策转向",
        12: "寻找成熟管理系统底座·若依进入视野",
    })
    notes(s[3], "真实经历：一开始想让AI全生成，发现基础框架消耗token大且不是业务核心。")

    # 5 底座选型：若依 vs 其他
    fill_slide(s[4], {
        21: "管理系统底座选型",
        22: "RUOYI AS FOUNDATION",
        7: "若依优势",
        17: "权限菜单现成·微服务架构·Java8兼容·文档完善",
        9: "业务聚焦",
        15: "只需实现yianlian模块·双服务层设计·多线路同步",
        8: "token节省",
        16: "基础代码无需生成·AI聚焦业务逻辑·效率显著提升",
        18: "技术约束",
        20: "一人维护·Java8·必须Win+Mac·30303不可改",
    })
    notes(s[4], "若依不是「套个壳」，而是经过AI对比自研/升Boot3后的理性决策。")

    # 6 用户端演进：Web到客户端
    fill_slide(s[5], {
        9: "用户端技术演进",
        10: "FROM WEB TO NATIVE CLIENT",
        12: "第一版：Web端",
        11: "快速验证·跨平台·HTTP登录·统一选线功能",
        14: "安全问题",
        13: "抓包可见密码·HTTP明文传输·与零信任矛盾",
        16: "第二版：客户端",
        15: "TLS 1.3加密·证书Pin·Qt跨平台·Win/Mac原生",
        18: "关键决策",
        17: "Web端下线·只保留管理端·安全优先",
        20: "AI角色",
        19: "列出Web/Electron/Qt/WPF四选一对比表",
    })
    notes(s[5], "Web端不是失败，是MVP验证。安全问题驱动技术升级。")

    # 7 节首页：AI辅助决策
    fill_slide(s[6], {
        4: "AI辅助架构决策",
        5: "AI-ASSISTED DESIGN",
        6: "Qt跨平台\n国际化设计",
        7: "人定约束\nAI列方案",
    })
    notes(s[6], "进入人机协作的核心：AI不是执行工具，而是决策助手。")

    # 8 客户端技术选型
    fill_slide(s[7], {
        12: "客户端技术栈四选一",
        13: "QT CROSS-PLATFORM",
        8: "Electron",
        9: "Qt 6+QML",
        10: "WPF",
        11: "Avalonia",
        15: "AI建议",
        14: "Electron快·WPF只Win·Avalonia不熟C#·Qt原生+C++可迁移",
        17: "人决策",
        16: "选Qt·否决Electron(安全面大)·Win/Mac一套代码",
        19: "ADR记录",
        18: "ADR-001桌面端技术栈·四选一对比·决策理由",
        21: "协议选型",
        20: "ADR-002·否HTTP复用网关·TLS 1.3+Protobuf",
    })
    notes(s[7], "关键：AI列方案，人基于安全/维护/能力做决策，结果落ADR文档。")

    # 9 AI辅助跨平台设计
    fill_slide(s[8], {
        4: "AI辅助：跨平台与国际化",
        5: "CROSS-PLATFORM & I18N",
        9: "跨平台设计",
        8: "Win/Mac共享QML·CMake构建·条件编译·平台特性隔离",
        13: "国际化设计",
        12: "多语言支持·资源文件管理·动态切换·文案外部化",
        11: "AI角色",
        10: "代码示例·最佳实践·CMakeLists配置·Qt语言家",
        7: "学习曲线",
        6: "从Qt零基础到完整客户端·AI即时导师·边做边学",
    })
    notes(s[8], "AI的价值：陌生技术域(Qt/QML)的即时导师，不需要先系统学习。")

    # 10 传输协议与安全分级
    fill_slide(s[9], {
        27: "传输协议与安全分级",
        28: "SECURITY BY DESIGN",
        26: "协议决策",
        25: "桌面端TLS 1.3 :9443·管理端HTTP网关·本地30303不可改",
        22: "S0人工验收",
        21: "TLS配置·证书Pin·密钥不进日志·抓包密文",
        24: "S1/S2分级",
        23: "S1人审逻辑·S2脚本断言·客户端不进accept.py",
    })
    notes(s[9], "安全不是配好栈就完事，每个新功能都要验证。S0永远只人验。")

    # 11 节首页：六阶段工作法
    fill_slide(s[10], {
        4: "六阶段工作法",
        5: "SIX-PHASE METHOD",
        6: "方法沉淀\n可复用·可推广",
        7: "多项目验证\n4个系统已用",
    })
    notes(s[10], "方法论核心，也是最大沉淀。已在部门推广、多项目验证。")

    # 12 六阶段流程
    fill_slide(s[11], {
        1: "六阶段流程全景",
        2: "STRUCTURED WORKFLOW",
        8: "①人工拆解",
        7: "模块清单·熟悉度标注·安全级·边界确认",
        4: "②AI共创",
        3: "架构方案·多案对比·spec人批·未批不写码",
        10: "③④分流",
        9: "陌生域先Demo突破·熟悉域按参照生成·不另起炉灶",
        6: "⑤⑥门禁",
        5: "gate→deploy→accept·质量门禁.md·回修≤2轮",
    })
    notes(s[11], "为什么是六阶段？每个阶段解决什么问题？后面详细展开。")

    # 13 为什么是六阶段
    fill_slide(s[12], {
        20: "为什么是六阶段？",
        21: "WHY SIX PHASES",
        9: "①②设计",
        8: "人主导方向·AI提供方案·决策权在人·防止方向跑偏",
        10: "③④实现",
        11: "按熟悉度分流·陌生先突破·熟悉按规范·AI加速执行",
        12: "⑤⑥验证",
        14: "人工门禁·脚本断言·钉钉回修·超2轮停下",
        13: "核心价值",
        15: "可重复·可管控·可验证·可推广",
    })
    notes(s[12], "六阶段不是拍脑袋，是从GenlotVPN实践中提炼、在其他项目验证的方法。")

    # 14 人与AI的配合
    fill_slide(s[13], {
        17: "人与AI如何配合",
        18: "HUMAN-AI COLLABORATION",
        5: "人的职责",
        6: "拆解边界",
        7: "批准spec",
        8: "S0验收",
        9: "业务目标拆解·熟悉度标注·架构决策·安全审核",
        10: "AI的职责",
        11: "列方案·生成代码·排错·学习",
        12: "多案对比·代码初稿·报错分析·即时导师·不承担责任",
    })
    notes(s[13], "职责边界清晰：人定方向和质量，AI加速学习和执行。")

    # 15 节首页：多项目验证
    fill_slide(s[14], {
        4: "多项目验证",
        5: "VERIFIED ACROSS PROJECTS",
        6: "GenlotVPN\n山西热线·统一系统",
        7: "验票系统\n部门已推广",
    })
    notes(s[14], "方法不是纸上谈兵，已在4个项目验证、部门内推广使用。")

    # 16 已验证项目
    fill_slide(s[15], {
        20: "已验证的项目",
        21: "PROJECTS VALIDATED",
        19: "GenlotVPN",
        18: "统一VPN管理·后端+Web+Qt客户端·首个完整验证",
        15: "山西热线",
        14: "同一套六阶段流程·迁移pipeline到新仓库·验证可复用性",
        17: "山西统一系统",
        16: "再次验证方法可移植·适配不同技术栈·质量门禁可重入",
        13: "验票系统",
        12: "第四个项目·部门内已开始推广·SOP文档已分享",
        11: "",
    })
    notes(s[15], "不是只做了一个项目，而是验证了方法的可推广性。每个项目都有pipeline/work目录可核验。")

    # 17 推广与沉淀
    fill_slide(s[16], {
        11: "推广与沉淀",
        12: "METHODOLOGY SHARING",
        10: "部门分享",
        9: "SOP文档·工程实施手册·六阶段流程讲解·Rule/Hook模板",
        6: "迁移验证",
        5: "pipeline可搬迁·流水线搬迁文档·gate/deploy/accept命令统一",
        8: "团队采用",
        7: "部门内已开始使用·适配各自项目·验证可推广性",
        3: "持续优化",
        2: "ADR记录踩坑·质量门禁记录·回修经验沉淀",
        1: "",
    })
    notes(s[16], "不是个人工作法，而是可推广的团队方法。有文档、有模板、有验证。")

    # 18 关键数据与边界
    fill_slide(s[17], {
        0: "关键数据与边界",
        14: "METRICS & BOUNDARIES",
        6: "01",
        7: "02",
        8: "03",
        9: "04",
        (12, 1): "编制对比",
        (12, 0): "传统≥4人(后端+前端+桌面+测试)·现在1人+AI",
        (10, 1): "时间节点",
        (10, 0): "08-16: 11:14 gate→12:03 deploy→12:22 accept→13:21客户端",
        (11, 1): "边界说明",
        (11, 0): "S0只人验·客户端不自动验收·生产验证码保持开",
        (13, 1): "推广范围",
        (13, 0): "部门内推广·4个项目验证·非全公司自动化",
    })
    notes(s[17], "主动讲边界：不说倍数、不说已投产、不说全自动。只讲可核验的数据。")

    # 19 总结与展望(使用s[18])
    fill_slide(s[18], {
        2: "技术分享",
        6: "架构决策可追溯\n六阶段可推广\n多项目已验证",
        12: "GENLOT VPN",
        17: "可以提问",
    })
    notes(s[18], "核心价值：1)架构决策过程(ADR可查) 2)工作方法沉淀(SOP可用) 3)多项目验证(证据可核)")

    dest = OUT
    try:
        prs.save(str(OUT))
    except PermissionError:
        dest = OUT.with_name("分享-研发AI全栈交付-40分钟-新.pptx")
        prs.save(str(dest))
    tmp.unlink(missing_ok=True)
    print("wrote", dest, "slides", len(prs.slides))


if __name__ == "__main__":
    build()
